# -*- coding: utf-8 -*-
"""工程図連携のスプレッドシートのひな形（xlsx）を作る。

node sync/template/build-sheets.mjs から呼ばれる（Code.gs の定数を JSON で受け取る）。
  python3 build_sheets.py <constants.json> <出力フォルダ>

作るファイル（Google ドライブにアップロードして「Google スプレッドシートに変換」して使う）:
  工程図連携_管理者用.xlsx          … 『連携』『会社マスタ』『視点マスタ』『案件マスタ』『設定』『使い方』（管理者だけが見る）
  工程図_エンジニア入力シート.xlsx … 『入力シート』『記入ルール・説明』『選択肢』（日本語／ベトナム語の併記。制作メンバーに共有）

見出しの文字は Code.gs が列を探すときに使うので変えない。色・列幅・注記・プルダウンはここで付ける。
Google スプレッドシートに変換した後、管理者シートのメニュー「0. かんたん初期設定」を押すと、
チェックボックスなど xlsx では表せないものが整う（このひな形の見出し・色はそのまま残る）。
"""
import json
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

FONT = 'Arial'
AUTHOR = '工程図連携'
# 色（Google スプレッドシートに変換しても残る）
C_HEAD = 'dde5f0'       # 見出し（Code.gs の formatStaffInputTab_ と同じ）
C_HEAD_REQ = 'c3d2e8'   # 必須の見出し（少し濃い）
C_HINT = 'f3f6fb'       # 案内行
C_TABLE = 'f7f9fc'      # 記入例の表
C_TITLE = '1a4f8b'      # 見出し文字（青）
C_VI = '666666'         # ベトナム語（灰色）
C_INPUT = 'fff2cc'      # 人が入れる欄（黄色）
C_LINE = 'd0d7e2'       # 罫線
THIN = Side(style='thin', color=C_LINE)
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
BOTTOM = Border(bottom=Side(style='thin', color=C_TITLE))
INPUT_ROWS = 1000  # 入力行の数（プルダウン・罫線を付ける範囲）


def font(bold=False, size=10, color=None, italic=False):
    return Font(name=FONT, bold=bold, size=size, color=color, italic=italic)


def fill(color):
    return PatternFill('solid', start_color=color, end_color=color)


def note(text, width=360, height=120):
    c = Comment(text, AUTHOR)
    c.width, c.height = width, height
    return c


def set_widths(ws, widths):
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def write(ws, row, col, value, **style):
    cell = ws.cell(row=row, column=col, value=value)
    cell.font = style.get('font', font())
    if style.get('fill'):
        cell.fill = fill(style['fill'])
    if 'align' in style:
        cell.alignment = style['align']
    if 'border' in style:
        cell.border = style['border']
    if 'fmt' in style:
        cell.number_format = style['fmt']
    if 'note' in style and style['note']:
        cell.comment = note(style['note'])
    return cell


def list_validation(formula, allow_blank=True, error=None, prompt=None, strict=False):
    dv = DataValidation(type='list', formula1=formula, allow_blank=allow_blank, showDropDown=False)
    dv.errorStyle = 'stop' if strict else 'warning'
    if error:
        dv.error, dv.errorTitle = error, '入力 / Nhập liệu'
    if prompt:
        dv.prompt, dv.promptTitle, dv.showInputMessage = prompt, '入力 / Nhập liệu', True
    dv.showErrorMessage = True
    return dv


def inline_list(values):
    return '"' + ','.join(str(v) for v in values) + '"'


# ============ エンジニア入力シート ============
def build_staff(d, out_dir):
    wb = Workbook()
    ws = wb.active
    ws.title = d['staffTab']
    cols = d['staffColumns']
    n = len(cols)
    required = {'code', 'number', 'request', 'cut', 'white', 'color', 'other'}

    # 1行目：案内（日本語／ベトナム語）
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=n)
    write(ws, 1, 1, d['staffHint'], font=font(color='555555'), fill=C_HINT,
          align=Alignment(vertical='center', wrap_text=True))
    ws.row_dimensions[1].height = 30

    # 2行目：見出し（日本語＋改行＋ベトナム語。Code.gs はこの日本語で列を探す）
    for i, c in enumerate(cols, start=1):
        write(ws, 2, i, c['label'] + '\n' + c['vi'], font=font(bold=True),
              fill=C_HEAD_REQ if c['key'] in required else C_HEAD,
              align=Alignment(horizontal='center', vertical='center', wrap_text=True), border=BORDER, note=c['note'])
    ws.row_dimensions[2].height = 42
    ws.freeze_panes = 'A3'
    set_widths(ws, [c['width'] * 1.35 for c in cols])

    # 選択肢タブ（プルダウンの元。視点は 100 件を超えるので、セル範囲で参照する）
    lists = wb.create_sheet('選択肢')
    lists.sheet_properties.tabColor = '999999'
    lists_def = [
        ('会社コード\nMã công ty', d['companyCodes']),
        ('新規or修正\nMới / Sửa', d['requestOptions']),
        ('視点\nGóc nhìn', d['viewOptions']),
        ('パターン\nPhương án', d['patternOptions']),
        ('時間（h）\nSố giờ', [float(h) if '.' in h else int(h) for h in d['hourOptions']]),
    ]
    lists.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(lists_def))
    write(lists, 1, 1, 'このタブは『' + d['staffTab'] + '』のプルダウンの選択肢です。手で変えないでください（管理者がマスタから更新します）。\n'
          'Tab này là danh sách lựa chọn cho tab『' + d['staffTab'] + '』. Vui lòng không chỉnh sửa (quản lý sẽ cập nhật từ master).',
          font=font(color='555555'), fill=C_HINT, align=Alignment(vertical='center', wrap_text=True))
    lists.row_dimensions[1].height = 52
    ranges = {}
    for i, (head, values) in enumerate(lists_def, start=1):
        write(lists, 2, i, head, font=font(bold=True), fill=C_HEAD, align=Alignment(horizontal='center', vertical='center', wrap_text=True), border=BORDER)
        for r, v in enumerate(values, start=3):
            write(lists, r, i, v, align=Alignment(horizontal='center'), fmt='0.##' if i == 5 else 'General')
        col = get_column_letter(i)
        ranges[i] = "'選択肢'!$%s$3:$%s$%d" % (col, col, 2 + len(values))
    lists.row_dimensions[2].height = 32
    set_widths(lists, [16, 16, 12, 12, 12])
    lists.freeze_panes = 'A3'

    # 3行目から：プルダウン・書式・罫線
    first, last = 3, 2 + INPUT_ROWS
    dvs = {
        'code': list_validation(ranges[1], prompt='リストから選ぶ（無い会社は管理者へ）\nChọn từ danh sách (nếu không có, báo quản lý)',
                                error='リストに無い会社コードです。管理者に確認してください。\nMã công ty không có trong danh sách. Hãy hỏi quản lý.'),
        'request': list_validation(ranges[2], strict=True, prompt='初回依頼は「新規」、2回目以降は「修正」\nLần đầu: Mới, từ lần 2: Sửa',
                                   error='リストから選んでください。\nHãy chọn từ danh sách.'),
        'cut': list_validation(ranges[3], prompt='リストから選ぶ（一覧は『' + d['staffRulesTab'] + '』）\nChọn mã góc nhìn từ danh sách',
                               error='リストに無い視点コードです（管理者が確認します）。\nMã góc nhìn không có trong danh sách (quản lý sẽ kiểm tra).'),
        'pattern': list_validation(ranges[4], prompt='パターン違いが無ければ空欄\nKhông có phương án khác thì để trống',
                                   error='A〜F から選んでください。\nHãy chọn A–F.'),
        'hours': list_validation(ranges[5], prompt='時間（h）。0.5 刻み。リストに無い数字も入れられます\nSố giờ (bước 0.5). Có thể nhập số khác',
                                 error='時間（h）を数字で入れてください。\nHãy nhập số giờ.'),
    }
    number_dv = DataValidation(type='whole', operator='between', formula1='1', formula2='99999', allow_blank=True)
    number_dv.errorStyle, number_dv.showErrorMessage = 'stop', True
    number_dv.errorTitle, number_dv.error = '案件番号 / Số dự án', '数字だけを入力してください（例 34）\nChỉ nhập số (VD: 34)'
    number_dv.promptTitle, number_dv.prompt, number_dv.showInputMessage = '案件番号 / Số dự án', '数字だけ（例 34）\nChỉ nhập số (VD: 34)', True
    for dv in list(dvs.values()) + [number_dv]:
        ws.add_data_validation(dv)
    center = Alignment(horizontal='center', vertical='center')
    left = Alignment(vertical='center')
    for i, c in enumerate(cols, start=1):
        letter = get_column_letter(i)
        rng = '%s%d:%s%d' % (letter, first, letter, last)
        key = c['key']
        if key == 'number':
            number_dv.add(rng)
        elif key in ('white', 'color', 'other'):
            dvs['hours'].add(rng)
        elif key in dvs:
            dvs[key].add(rng)
        for r in range(first, last + 1):
            cell = ws.cell(row=r, column=i)
            cell.font = font()
            cell.border = BORDER
            cell.alignment = center if key in ('code', 'number', 'request', 'cut', 'pattern', 'white', 'color', 'other') else left
            if c.get('hours'):
                cell.number_format = '0.##'
            elif key == 'number':
                cell.number_format = '0'
    ws.sheet_properties.tabColor = C_TITLE
    ws.sheet_view.zoomScale = 100

    build_staff_rules(wb, d)
    # タブの順：入力シート → 記入ルール/説明 → 選択肢
    wb.move_sheet('選択肢', offset=1)
    path = Path(out_dir) / '工程図_エンジニア入力シート.xlsx'
    wb.save(path)
    return path


def build_staff_rules(wb, d):
    """『記入ルール/説明』：日本語の下にベトナム語。視点コード表と記入例は表にする（本文は B 列から。A 列は余白）"""
    # Excel はタブ名に「/」を使えないので『記入ルール・説明』で作る（管理者シートの「0. かんたん初期設定」が『記入ルール/説明』に改名する）
    ws = wb.create_sheet(d['staffRulesTabXlsx'])
    ws.sheet_properties.tabColor = '5b8def'
    cols = d['staffColumns']
    width = len(cols)
    r = 2
    for kind, ja, vi in d['staffRulesLines']:
        if kind == 'title':
            write(ws, r, 2, ja, font=font(bold=True, size=14, color=C_TITLE))
            write(ws, r + 1, 2, vi, font=font(size=12, color=C_VI))
            r += 3
        elif kind == 'h':
            if r > 5:
                r += 1
            for c in range(2, width + 2):
                ws.cell(row=r, column=c).border = BOTTOM
            write(ws, r, 2, ja + ' ／ ' + vi.replace('■', '').strip(), font=font(bold=True, size=11, color=C_TITLE), border=BOTTOM)
            r += 1
        else:
            write(ws, r, 2, ja, font=font())
            write(ws, r + 1, 2, '   ' + vi, font=font(color=C_VI))
            r += 2
    # 視点コード表（コード／日本語／ベトナム語）。STAFF_RULES_LINES の最後が「■ 視点コード」なので、その直後に続ける
    head = Alignment(horizontal='center', vertical='center', wrap_text=True)
    for c, text in enumerate(['コード\nMã', '日本語', 'Tiếng Việt'], start=2):
        write(ws, r, c, text, font=font(bold=True), fill=C_HEAD, align=head, border=BORDER)
    ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=5)
    ws.merge_cells(start_row=r, start_column=6, end_row=r, end_column=9)
    ws.row_dimensions[r].height = 30
    r += 1
    for v in d['viewCodeTable']:
        write(ws, r, 2, v['code'], font=font(bold=True), align=Alignment(horizontal='center'), border=BORDER, fill=C_TABLE)
        write(ws, r, 3, v['ja'], border=BORDER, fill=C_TABLE)
        write(ws, r, 6, v['vi'], font=font(color=C_VI), border=BORDER, fill=C_TABLE)
        for c in (4, 5, 7, 8, 9):
            ws.cell(row=r, column=c).border = BORDER
            ws.cell(row=r, column=c).fill = fill(C_TABLE)
        ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=5)
        ws.merge_cells(start_row=r, start_column=6, end_row=r, end_column=9)
        r += 1
    write(ws, r, 2, '※ 末尾の数字＝①②…（外観目線視点② → EX2）。表に無い視点は管理者に連絡 ／ Số cuối = ①②… (VD: ngoại thất tầm mắt ② → EX2). Nếu không có trong bảng, liên hệ quản lý',
          font=font(color=C_VI, italic=True))
    r += 2
    # 記入例
    for c in range(2, width + 2):
        ws.cell(row=r, column=c).border = BOTTOM
    write(ws, r, 2, "■ 記入例（読むだけ。実際の入力は『%s』タブへ） ／ Ví dụ (chỉ để xem, nhập thật ở tab『%s』)" % (d['staffTab'], d['staffTab']),
          font=font(bold=True, size=11, color=C_TITLE), border=BOTTOM)
    r += 1
    for i, c in enumerate(cols, start=2):
        write(ws, r, i, c['label'] + '\n' + c['vi'], font=font(bold=True), fill=C_HEAD, align=head, border=BORDER)
    ws.row_dimensions[r].height = 42
    r += 1
    for ex in d['staffRulesExamples']:
        for i, v in enumerate(ex, start=2):
            key = cols[i - 2]['key']
            write(ws, r, i, v, fill=C_TABLE, border=BORDER,
                  align=Alignment(horizontal='center') if key in ('code', 'number', 'request', 'cut', 'pattern', 'white', 'color', 'other') else Alignment(),
                  fmt='0.##' if cols[i - 2].get('hours') else 'General')
        r += 1
    set_widths(ws, [3] + [c['width'] * 1.35 for c in cols])


# ============ 管理者シート（工程図連携） ============
LINK_NOTES = {
    'key': '取込キー：案件コード::視点名（パターン込み）::n回目。自動で入ります（変えない）',
    'status': '未送信（自動登録を待っている）／要確認／更新あり／登録済み／エラー／元シートから消えた',
    'send': 'チェックすると次の自動実行（またはメニュー「2. 工程図へ送信」）で工程図に登録します。「0. かんたん初期設定」でチェックボックスになります',
    'inputDate': '最初に転記した日（自動。あとで変わりません）',
    'code': '会社コード＋案件番号（例 RIC.34）。自動',
    'name': '工程図の案件名。自動判定の結果が入り、人が直せます（入力シートの社外案件名が書き換わったときは追従）',
    'cut': '視点コード（例 EX1）。自動',
    'pattern': 'パターン（A・B…）。自動',
    'extName': 'お客様向けの視点名（例 外観目線視点①_パターンA）。自動判定の結果が入り、人が直せます',
    'round': '同じ案件＋視点の何回目か。自動',
    'request': '入力シートの「新規or修正」。自動',
    'deadline': '納期（任意）。登録前に入れたものが工程図に入ります',
    'white': 'ホワイトパースまでの時間（h）。自動',
    'color': 'カラーパースまでの時間（h）。自動',
    'other': '写真合成・人物など その他の時間（h）。自動',
    'item': '制作項目（旧レイアウト用。今の入力シートには無いので、人が入れた値がそのまま残ります）',
    'note': '入力シートの「メモ」。自動',
    'company': '工程図の会社名（顧客マスタと同じ表記）。自動判定の結果が入り、人が直せます',
    'contact': 'お客様担当者（任意）。『案件マスタ』から入ります',
    'category': '区分（外観目線・内観目線…）。『視点マスタ』から自動判定。人が直せます',
    'kind': '種類（パース／写真合成／モデル／VR／動画）。『視点マスタ』から自動判定。人が直せます',
    'stepKind': '新規／追加／修正（無料）／変更（有料）。入力シートの新規or修正から入ります。有料の変更・追加は人が直す',
    'assignee': '社内担当者（任意）。空なら『設定』の「既定の担当者」',
    'memo': '工程図に入れるメモ（任意）',
    'exclude': 'チェックすると工程図に送りません（二重入力など）。「0. かんたん初期設定」でチェックボックスになります',
    'transferredAt': '最後に内容が変わった日時（自動）',
    'sentAt': '工程図に送った日時（自動）',
    'result': '自動判定・送信の結果（自動。「要確認」の理由もここ）',
    'srcRow': '入力シートの行番号（自動）',
    'link': '入力シートのサーバーリンク（自動）',
}
LINK_WIDTHS = {
    'key': 26, 'status': 11, 'send': 9, 'inputDate': 11, 'code': 11, 'name': 20, 'cut': 8, 'pattern': 9, 'extName': 24, 'round': 5,
    'request': 11, 'deadline': 11, 'white': 10, 'color': 10, 'other': 10, 'item': 10, 'note': 24, 'company': 22, 'contact': 13,
    'category': 11, 'kind': 10, 'stepKind': 13, 'assignee': 11, 'memo': 22, 'exclude': 8, 'transferredAt': 17, 'sentAt': 17,
    'result': 40, 'srcRow': 9, 'link': 32,
}
LINK_FORMATS = {'inputDate': 'yyyy-mm-dd', 'deadline': 'yyyy-mm-dd', 'transferredAt': 'yyyy-mm-dd hh:mm', 'sentAt': 'yyyy-mm-dd hh:mm',
                'white': '0.##', 'color': '0.##', 'other': '0.##', 'round': '0', 'srcRow': '0'}


def build_admin(d, out_dir):
    wb = Workbook()
    colors = d['linkHeaderColors']
    auto_fill, human_fill = colors['auto'].lstrip('#'), colors['human'].lstrip('#')
    head = Alignment(horizontal='center', vertical='center', wrap_text=True)

    # 『連携』
    ws = wb.active
    ws.title = d['sheet']['LINK']
    ws.sheet_properties.tabColor = C_TITLE
    headers = d['linkHeaders']
    for i, h in enumerate(headers, start=1):
        write(ws, 1, i, h['label'], font=font(bold=True), fill=human_fill if h['human'] else auto_fill, align=head, border=BORDER,
              note=LINK_NOTES.get(h['key']))
    ws.row_dimensions[1].height = 30
    ws.freeze_panes = 'F2'
    set_widths(ws, [LINK_WIDTHS.get(h['key'], 12) for h in headers])
    last = 1 + INPUT_ROWS
    dvs = {
        'status': list_validation(inline_list(d['statusOptions'])),
        'category': list_validation(inline_list(d['linkChoices']['categories'])),
        'kind': list_validation(inline_list(d['linkChoices']['kinds'])),
        'stepKind': list_validation(inline_list(d['stepKindOptions'])),
    }
    for dv in dvs.values():
        ws.add_data_validation(dv)
    for i, h in enumerate(headers, start=1):
        letter = get_column_letter(i)
        if h['key'] in dvs:
            dvs[h['key']].add('%s2:%s%d' % (letter, letter, last))
        fmt = LINK_FORMATS.get(h['key'])
        if fmt:
            for r in range(2, last + 1):
                ws.cell(row=r, column=i).number_format = fmt
        if h['key'] in ('status', 'send', 'exclude', 'round', 'request', 'cut', 'pattern', 'inputDate'):
            for r in range(2, last + 1):
                ws.cell(row=r, column=i).alignment = Alignment(horizontal='center')

    # 『会社マスタ』
    cm = wb.create_sheet(d['sheet']['COMPANY'])
    heads = d['companyMaster']['headers']
    notes = ['エンジニアが入力シートで選ぶ会社コード（案件コードの英字部分。例 RIC）',
             '工程図の顧客マスタと同じ表記にする（メニュー「工程図の会社名一覧を取り込む」で E 列に一覧が出ます）',
             '備考（「表記ゆれ」と書いた行は初期設定で「入力シートに出さない」にチェックが入ります）',
             'チェックした会社コードは、エンジニア入力シートの「会社コード」のプルダウンに出しません。変えたらメニュー「エンジニア入力シートを整える」']
    for i, h in enumerate(heads, start=1):
        write(cm, 1, i, h, font=font(bold=True), fill=C_HEAD, align=head, border=BORDER, note=notes[i - 1])
    for r, row in enumerate(d['companyMaster']['rows'], start=2):
        for i, v in enumerate(row, start=1):
            write(cm, r, i, v, border=BORDER, align=Alignment(horizontal='center') if i in (1, 4) else Alignment())
    cm.row_dimensions[1].height = 30
    cm.freeze_panes = 'A2'
    set_widths(cm, [22, 30, 26, 20, 30])

    # 『視点マスタ』
    vm = wb.create_sheet(d['sheet']['VIEW'])
    heads = d['viewMaster']['headers']
    notes = ['エンジニアが書く視点コードの英字部分（EX1 → EX）。足すと「エンジニア入力シートを整える」でプルダウンにも出ます',
             '工程図の視点の「区分」（売上の制作種類）。VR・動画は空でよい',
             'パース／写真合成／モデル／VR／動画。新しい名前を書くと、その名前の1ステップで登録',
             d['viewMaster']['externalNote'], '']
    for i, h in enumerate(heads, start=1):
        write(vm, 1, i, h, font=font(bold=True), fill=C_HEAD, align=head, border=BORDER, note=notes[i - 1] or None)
    for r, row in enumerate(d['viewMaster']['rows'], start=2):
        for i, v in enumerate(row, start=1):
            write(vm, r, i, v, border=BORDER, align=Alignment(horizontal='center') if i == 1 else Alignment())
    vm.row_dimensions[1].height = 30
    vm.freeze_panes = 'A2'
    set_widths(vm, [22, 14, 12, 16, 48])

    # 『案件マスタ』（登録は任意。ここに書いた案件は「要確認」になりにくい）
    pm = wb.create_sheet(d['sheet']['PROJECT'])
    notes = ['スタッフが書く社内案件名（例 RIC.34）。大文字小文字・点・空白は無視して照合', 'お客様向けの案件名（工程図の案件名になる）',
             '工程図の会社名（空なら会社マスタで判定）', '工程図のお客様担当者（任意）', '']
    for i, h in enumerate(d['projectMasterHeaders'], start=1):
        write(pm, 1, i, h, font=font(bold=True), fill=C_HEAD, align=head, border=BORDER, note=notes[i - 1] or None)
    pm.row_dimensions[1].height = 30
    pm.freeze_panes = 'A2'
    set_widths(pm, [22, 32, 22, 16, 30])
    for r in range(2, 202):
        for i in range(1, 6):
            pm.cell(row=r, column=i).border = BORDER

    # 『設定』
    st = wb.create_sheet(d['sheet']['SETTINGS'])
    for i, h in enumerate(['項目', '値', '説明'], start=1):
        write(st, 1, i, h, font=font(bold=True), fill=C_HEAD, align=head, border=BORDER)
    auto_dv = None
    for r, s in enumerate(d['settings'], start=2):
        write(st, r, 1, s['label'], font=font(bold=True), border=BORDER, align=Alignment(vertical='center'))
        write(st, r, 2, s['value'], border=BORDER, fill=C_INPUT if s['key'] in ('fileId', 'autoSend', 'autoWait', 'defaultAssignee') else None,
              align=Alignment(vertical='center'))
        write(st, r, 3, s['note'], font=font(color='555555'), border=BORDER, align=Alignment(wrap_text=True, vertical='center'))
        if s['key'] == 'autoSend':
            auto_dv = list_validation('"する,しない"')
            st.add_data_validation(auto_dv)
            auto_dv.add('B%d' % r)
        if s['key'] == 'fileId':
            st.cell(row=r, column=2).comment = note('「工程図 エンジニア入力シート」を新しく作った場合は、その URL の /d/ と /edit の間の文字列に差し替える')
    set_widths(st, [30, 46, 96])
    st.freeze_panes = 'A2'

    # 『使い方』
    hw = wb.create_sheet(d['sheet']['HOWTO'])
    for r, line in enumerate(d['howtoLines'], start=1):
        if r == 1:
            write(hw, r, 1, line, font=font(bold=True, size=14, color=C_TITLE))
        elif line.startswith('■'):
            write(hw, r, 1, line, font=font(bold=True, size=11, color=C_TITLE))
        else:
            write(hw, r, 1, line)
    set_widths(hw, [130])

    path = Path(out_dir) / '工程図連携_管理者用.xlsx'
    wb.save(path)
    return path


def main():
    d = json.loads(Path(sys.argv[1]).read_text(encoding='utf-8'))
    out_dir = Path(sys.argv[2])
    out_dir.mkdir(parents=True, exist_ok=True)
    for p in (build_admin(d, out_dir), build_staff(d, out_dir)):
        print('wrote', p)


if __name__ == '__main__':
    main()
